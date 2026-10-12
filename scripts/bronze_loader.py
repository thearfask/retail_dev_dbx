import hashlib
import json
import shutil
import uuid
import zipfile

from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
import requests
import yaml

from delta.tables import DeltaTable
from pyspark.sql import functions as F
from pyspark.sql import Window


def now():
    return datetime.now(timezone.utc)


def checksum(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_config(spark, path):
    raw = Path(path).read_bytes()
    config = yaml.safe_load(raw)

    config["checksum"] = hashlib.sha256(raw).hexdigest()

    env = config["environment"]
    config["namespace"] = f"{env['catalog']}.{env['schema']}"
    config["root"] = (
        f"/Volumes/{env['catalog']}/{env['schema']}/{env['volume']}"
    )

    config["parsed_schemas"] = {
        name: spark.createDataFrame([], ddl).schema
        for name, ddl in config["schemas"].items()
    }

    return config


def table(config, name):
    return f"{config['namespace']}.{name}"


def setup(spark, config, reset=False):
    """
    Creates project tables.
    Optional reset removes only tables used by this implementation.
    """
    namespace = config["namespace"]
    env = config["environment"]

    if reset and env["schema"] != "retail_dev_schema":
        raise ValueError("Reset is allowed only in retail_dev_schema")

    spark.sql(f"CREATE SCHEMA IF NOT EXISTS {namespace}")
    spark.sql(
        f"CREATE VOLUME IF NOT EXISTS {namespace}.{env['volume']}"
    )

    mapping = {
        "bronze": "bronze_transactions",
        "quarantine": "retail_quarantine",
        "monitoring": "retail_monitoring",
        "silver_order_items": "silver_order_items",
        "silver_customers_current": "silver_customers_current",
        "silver_customers_history": "silver_customers_history",
    }

    for schema_name, table_name in mapping.items():
        target = table(config, table_name)

        if reset:
            spark.sql(f"DROP TABLE IF EXISTS {target}")

        spark.sql(
            f"CREATE TABLE IF NOT EXISTS {target} "
            f"({config['schemas'][schema_name]}) USING DELTA"
        )

        existing = spark.table(target).schema
        approved = config["parsed_schemas"][schema_name]

        old_types = {
            field.name: field.dataType.json()
            for field in existing
        }
        new_types = {
            field.name: field.dataType.json()
            for field in approved
        }

        # Reject incompatible schemas instead of silently changing data.
        if old_types != new_types:
            raise ValueError(
                f"Schema mismatch: {target}. "
                "Use an intentional dev reset or design a migration."
            )


def prepare_source(config):
    """
    Download original UCI workbook once and convert to JSONL.

    Workbook conversion runs on the driver in read-only streaming mode.
    All subsequent transformations use Spark.
    """
    root = Path(config["root"]) / "source" / "online_retail_ii"
    root.mkdir(parents=True, exist_ok=True)

    archive = root / "source.zip"
    workbook_path = root / config["source"]["workbook"]
    output = root / "transactions.jsonl"
    manifest_path = root / "manifest.json"

    if manifest_path.exists() and output.exists():
        manifest = json.loads(
            manifest_path.read_text(encoding="utf-8")
        )

        if checksum(output) != manifest["json_checksum"]:
            raise RuntimeError("Prepared source checksum changed")

        return output, manifest

    if not archive.exists():
        temporary = root / "source.zip.part"

        with requests.get(
            config["source"]["url"],
            stream=True,
            timeout=(10, 120),
        ) as response:
            response.raise_for_status()

            with temporary.open("wb") as handle:
                for chunk in response.iter_content(1024 * 1024):
                    if chunk:
                        handle.write(chunk)

        # Validate ZIP before publishing it.
        with zipfile.ZipFile(temporary) as zipped:
            if zipped.testzip() is not None:
                raise RuntimeError("Downloaded ZIP is corrupt")

        temporary.replace(archive)

    if not workbook_path.exists():
        with zipfile.ZipFile(archive) as zipped:
            members = [
                name for name in zipped.namelist()
                if Path(name).name == config["source"]["workbook"]
            ]

            if len(members) != 1:
                raise RuntimeError("Expected workbook not found")

            with (
                zipped.open(members[0]) as incoming,
                workbook_path.open("wb") as outgoing,
            ):
                shutil.copyfileobj(incoming, outgoing)

    source_hash = checksum(workbook_path)
    temporary_output = root / "transactions.jsonl.part"
    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    row_count = 0

    def normalize_id(value):
        if value is None:
            return None
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    try:
        with temporary_output.open(
            "w", encoding="utf-8"
        ) as handle:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = [
                    str(value).strip().lower()
                    for value in next(rows)
                ]

                def find(*names):
                    for name in names:
                        if name.lower() in headers:
                            return headers.index(name.lower())
                    raise ValueError(
                        f"Missing column {names} in {sheet.title}"
                    )

                positions = {
                    "invoice": find("Invoice", "InvoiceNo"),
                    "stock_code": find("StockCode"),
                    "description": find("Description"),
                    "quantity": find("Quantity"),
                    "invoice_date": find("InvoiceDate"),
                    "unit_price": find("Price", "UnitPrice"),
                    "customer_id": find("Customer ID", "CustomerID"),
                    "country": find("Country"),
                }

                for row_number, values in enumerate(rows, start=2):
                    if all(value is None for value in values):
                        continue

                    record = {}

                    for name, index in positions.items():
                        value = values[index]

                        if name == "invoice_date":
                            record[name] = (
                                value.isoformat()
                                if isinstance(value, datetime)
                                else None if value is None
                                else str(value)
                            )
                        elif name in {
                            "invoice", "stock_code", "customer_id"
                        }:
                            record[name] = normalize_id(value)
                        else:
                            record[name] = (
                                None if value is None else str(value)
                            )

                    # Invoice + product is not guaranteed unique.
                    identity = (
                        f"{source_hash}:{sheet.title}:{row_number}"
                    )
                    record["record_id"] = hashlib.sha256(
                        identity.encode()
                    ).hexdigest()

                    record["source_sheet"] = sheet.title
                    record["source_row"] = row_number

                    handle.write(
                        json.dumps(record, ensure_ascii=False) + "\n"
                    )
                    row_count += 1

    finally:
        workbook.close()

    temporary_output.replace(output)

    manifest = {
        "source_url": config["source"]["url"],
        "workbook_checksum": source_hash,
        "json_checksum": checksum(output),
        "row_count": row_count,
        "prepared_at": now().isoformat(),
    }

    manifest_path.write_text(
        json.dumps(manifest, indent=2),
        encoding="utf-8",
    )

    return output, manifest


def audit(spark, config, batch_id, run_id, **changes):
    target = table(config, "retail_monitoring")
    schema = spark.table(target).schema

    keys = {
        "batch_id": batch_id,
        "ingestion_run_id": run_id,
    }

    unknown = set(changes) - set(schema.fieldNames())
    if unknown:
        raise ValueError(f"Unknown audit fields: {unknown}")

    values = {**keys, **changes}
    row = {field.name: values.get(field.name) for field in schema}
    incoming = spark.createDataFrame([row], schema)

    (
        DeltaTable.forName(spark, target)
        .alias("t")
        .merge(
            incoming.alias("s"),
            """
            t.batch_id = s.batch_id
            AND t.ingestion_run_id = s.ingestion_run_id
            """,
        )
        .whenMatchedUpdate(
            set={
                name: f"s.`{name}`"
                for name in changes
            }
        )
        .whenNotMatchedInsertAll()
        .execute()
    )


def read_source(spark, config, source_path):
    return (
        spark.read.text(str(source_path))
        .select(F.col("value").alias("raw_record"))
        .withColumn(
            "record",
            F.from_json(
                "raw_record",
                config["parsed_schemas"]["source"],
                {"mode": "FAILFAST"},
            ),
        )
        .withColumn(
            "source_date",
            F.expr(
                "CAST(try_cast(record.invoice_date "
                "AS TIMESTAMP_NTZ) AS DATE)"
            ),
        )
    )


def ingest_range(
    spark,
    config,
    source_path,
    source_manifest,
    start,
    end,
    load_type="HISTORICAL",
    trigger_type="MANUAL",
    refresh=False,
):
    """
    Process [start, end).
    Same range and source contract skip when target counts agree.
    """
    start_date = date.fromisoformat(start)
    end_date = date.fromisoformat(end)

    if start_date >= end_date:
        raise ValueError("Start must be before end")

    run_id = str(uuid.uuid4())
    batch_id = f"batch_retail_{start_date:%Y%m%d}_{end_date:%Y%m%d}"

    monitoring = table(config, "retail_monitoring")
    bronze_table = table(config, "bronze_transactions")

    previous_rows = (
        spark.table(monitoring)
        .filter(
            (F.col("batch_id") == batch_id)
            & (F.col("status") == "SUCCEEDED")
        )
        .orderBy(F.col("completed_at").desc())
        .limit(1)
        .collect()
    )

    previous = previous_rows[0].asDict() if previous_rows else None

    audit(
        spark, config, batch_id, run_id,
        trigger_type=trigger_type,
        load_type=load_type,
        status="RUNNING",
        window_start=start_date,
        window_end=end_date,
        source_uri=config["source"]["url"],
        source_checksum=source_manifest["json_checksum"],
        contract_checksum=config["checksum"],
        started_at=now(),
    )

    try:
        # Range identity protects replay, not arbitrary target edits.
        if previous and not refresh:
            same_input = (
                previous["source_checksum"]
                == source_manifest["json_checksum"]
                and previous["contract_checksum"] == config["checksum"]
            )

            target_count = (
                spark.table(bronze_table)
                .filter(F.col("batch_id") == batch_id)
                .count()
            )

            if same_input and target_count == previous["accepted_count"]:
                audit(
                    spark, config, batch_id, run_id,
                    status="SKIPPED",
                    expected_count=previous["expected_count"],
                    accepted_count=0,
                    rejected_count=0,
                    completed_at=now(),
                    original_success_run_id=previous[
                        "ingestion_run_id"
                    ],
                )
                return run_id, "SKIPPED"

            if not same_input:
                raise RuntimeError(
                    "Source or contract changed; explicitly use refresh=True"
                )

        source = read_source(spark, config, source_path)

        # Fail rather than silently dropping records without usable dates.
        if source.filter(F.col("source_date").isNull()).limit(1).count():
            raise RuntimeError(
                "Source has unparseable dates; date routing needs correction"
            )

        selected = source.filter(
            (F.col("source_date") >= F.lit(start_date))
            & (F.col("source_date") < F.lit(end_date))
        )

        reason = (
            F.when(
                F.col("record.record_id").isNull(),
                F.lit("MISSING_SOURCE_ROW_ID"),
            )
            .when(
                F.col("record.invoice").isNull(),
                F.lit("MISSING_INVOICE"),
            )
            .when(
                F.col("record.stock_code").isNull(),
                F.lit("MISSING_STOCK_CODE"),
            )
        )

        selected = selected.withColumn("rejection_reason", reason)
        expected_count = selected.count()

        rejects = selected.filter(
            F.col("rejection_reason").isNotNull()
        )
        rejected_count = rejects.count()

        quarantine = rejects.select(
            F.col("record.record_id").alias("record_id"),
            "raw_record",
            F.lit(batch_id).alias("batch_id"),
            F.lit(run_id).alias("ingestion_run_id"),
            "rejection_reason",
            F.current_timestamp().alias("rejected_at"),
        )

        (
            quarantine.write.format("delta")
            .mode("overwrite")
            .option("replaceWhere", f"batch_id = '{batch_id}'")
            .saveAsTable(table(config, "retail_quarantine"))
        )

        if rejected_count:
            raise RuntimeError(
                f"Strict Bronze gate rejected {rejected_count} rows"
            )

        accepted = selected.select(
            F.col("record.record_id").alias("record_id"),
            "raw_record",
            "source_date",
            F.lit(batch_id).alias("batch_id"),
            F.lit(run_id).alias("ingestion_run_id"),
            F.lit(config["source"]["url"]).alias("source_uri"),
            F.lit(str(source_path)).alias("source_file"),
            F.lit(source_manifest["json_checksum"]).alias(
                "source_checksum"
            ),
            F.lit(config["contract_version"]).alias(
                "contract_version"
            ),
            F.lit(config["checksum"]).alias("contract_checksum"),
            F.current_timestamp().alias("ingested_at"),
        )

        (
            accepted.write.format("delta")
            .mode("overwrite")
            .option(
                "replaceWhere",
                f"source_date >= '{start}' AND source_date < '{end}'",
            )
            .saveAsTable(bronze_table)
        )

        actual = (
            spark.table(bronze_table)
            .filter(
                (F.col("source_date") >= F.lit(start_date))
                & (F.col("source_date") < F.lit(end_date))
            )
            .count()
        )

        if actual != expected_count:
            raise RuntimeError("Bronze reconciliation failed")

        audit(
            spark, config, batch_id, run_id,
            status="SUCCEEDED",
            expected_count=expected_count,
            accepted_count=actual,
            rejected_count=0,
            completed_at=now(),
        )

        return run_id, "SUCCEEDED"

    except Exception as error:
        audit(
            spark, config, batch_id, run_id,
            status="FAILED",
            error_message=str(error)[:4000],
            completed_at=now(),
        )
        raise


def build_silver_and_gold(spark, config):
    """
    Deterministic full Silver rebuild for this bounded practice dataset.
    Bronze ingestion remains range-based.

    Customer SCD2 records observed snapshot changes, not reconstructed
    historical customer changes.
    """
    namespace = config["namespace"]

    bronze = spark.table(table(config, "bronze_transactions"))

    typed = (
        bronze.withColumn(
            "record",
            F.from_json(
                "raw_record",
                config["parsed_schemas"]["source"],
                {"mode": "FAILFAST"},
            ),
        )
        .select(
            "record.*",
            "source_date",
            "batch_id",
            "ingestion_run_id",
        )
        .withColumn(
            "invoice_at",
            F.expr("try_cast(invoice_date AS TIMESTAMP_NTZ)"),
        )
        .withColumn(
            "quantity",
            F.expr("try_cast(quantity AS BIGINT)"),
        )
        .withColumn(
            "unit_price",
            F.expr("try_cast(unit_price AS DECIMAL(18,4))"),
        )
        .withColumn("country", F.trim("country"))
        .withColumn("description", F.trim("description"))
    )

    # Missing customer IDs are valid anonymous sales.
    invalid = typed.filter(
        F.col("invoice_at").isNull()
        | F.col("quantity").isNull()
        | F.col("unit_price").isNull()
    )

    if invalid.limit(1).count():
        raise RuntimeError(
            "Silver numeric/date validation failed; investigate before publishing"
        )

    items = (
        typed.withColumn(
            "is_cancellation",
            F.upper("invoice").startswith("C"),
        )
        .withColumn(
            "line_amount",
            (F.col("quantity") * F.col("unit_price")).cast(
                "decimal(28,4)"
            ),
        )
        .select(
            "record_id", "invoice", "stock_code",
            "description", "quantity", "invoice_at",
            "unit_price", "customer_id", "country",
            "source_date", "batch_id", "ingestion_run_id",
            "is_cancellation", "line_amount",
        )
    )

    if (
        items.groupBy("record_id").count()
        .filter("count > 1").limit(1).count()
    ):
        raise RuntimeError("Duplicate source row IDs in Silver")

    (
        items.write.format("delta")
        .mode("overwrite")
        .saveAsTable(table(config, "silver_order_items"))
    )

    # Latest country observed in transactions; tie-break deterministically.
    window = Window.partitionBy("customer_id").orderBy(
        F.col("invoice_at").desc(),
        F.col("record_id").desc(),
    )

    customers = (
        items.filter(
            F.col("customer_id").isNotNull()
            & (F.trim("customer_id") != "")
        )
        .withColumn("rn", F.row_number().over(window))
        .filter("rn = 1")
        .select("customer_id", "country")
        .withColumn(
            "attribute_hash",
            F.sha2(
                F.to_json(
                    F.struct("country"),
                    {"ignoreNullFields": "false"},
                ),
                256,
            ),
        )
    )

    customers.createOrReplaceTempView("incoming_customers")

    # SCD1: latest observed customer country.
    spark.sql(f"""
    MERGE INTO {namespace}.silver_customers_current t
    USING incoming_customers s
    ON t.customer_id = s.customer_id
    WHEN MATCHED AND t.attribute_hash <> s.attribute_hash THEN
      UPDATE SET
        t.country = s.country,
        t.attribute_hash = s.attribute_hash,
        t.updated_at = current_timestamp()
    WHEN NOT MATCHED THEN
      INSERT (customer_id, country, attribute_hash, updated_at)
      VALUES (
        s.customer_id, s.country, s.attribute_hash,
        current_timestamp()
      )
    """)

    # SCD2 staging produces expiration + insertion for changed customers.
    spark.sql(f"""
    CREATE OR REPLACE TEMP VIEW customer_history_stage AS
    SELECT
        s.customer_id AS merge_key,
        s.customer_id,
        s.country,
        s.attribute_hash
    FROM incoming_customers s

    UNION ALL

    SELECT
        CAST(NULL AS STRING) AS merge_key,
        s.customer_id,
        s.country,
        s.attribute_hash
    FROM incoming_customers s
    JOIN {namespace}.silver_customers_history t
      ON s.customer_id = t.customer_id
     AND t.is_current = true
    WHERE s.attribute_hash <> t.attribute_hash
    """)

    spark.sql(f"""
    MERGE INTO {namespace}.silver_customers_history t
    USING customer_history_stage s
    ON t.customer_id = s.merge_key AND t.is_current = true

    WHEN MATCHED AND t.attribute_hash <> s.attribute_hash THEN
      UPDATE SET
        t.is_current = false,
        t.valid_to = current_timestamp()

    WHEN NOT MATCHED THEN
      INSERT (
        customer_id, country, attribute_hash,
        valid_from, valid_to, is_current
      )
      VALUES (
        s.customer_id, s.country, s.attribute_hash,
        current_timestamp(), NULL, true
      )
    """)

    spark.sql(f"""
    CREATE OR REPLACE TABLE {namespace}.silver_orders
    USING DELTA AS
    SELECT
      invoice,
      MIN(invoice_at) AS invoice_at,
      COUNT(*) AS line_count,
      SUM(quantity) AS units,
      SUM(line_amount) AS signed_amount,
      MAX(CAST(is_cancellation AS INT)) = 1 AS is_cancellation
    FROM {namespace}.silver_order_items
    GROUP BY invoice
    """)

    spark.sql(f"""
    CREATE OR REPLACE TABLE {namespace}.silver_products
    USING DELTA AS
    SELECT stock_code, description
    FROM (
      SELECT
        stock_code,
        description,
        ROW_NUMBER() OVER (
          PARTITION BY stock_code
          ORDER BY invoice_at DESC, record_id DESC
        ) AS rn
      FROM {namespace}.silver_order_items
    )
    WHERE rn = 1
    """)

    spark.sql(f"""
    CREATE OR REPLACE TABLE {namespace}.gold_daily_sales
    USING DELTA AS
    SELECT
      source_date,
      country,
      COUNT(*) AS line_count,
      COUNT(DISTINCT CASE
        WHEN NOT is_cancellation AND quantity > 0
        THEN invoice END
      ) AS sale_invoice_count,
      SUM(CASE
        WHEN quantity > 0 AND NOT is_cancellation
        THEN line_amount ELSE 0 END
      ) AS gross_sales_gbp,
      SUM(CASE
        WHEN quantity < 0
        THEN -line_amount ELSE 0 END
      ) AS negative_adjustments_gbp,
      SUM(line_amount) AS net_signed_value_gbp
    FROM {namespace}.silver_order_items
    GROUP BY source_date, country
    """)

    spark.sql(f"""
    CREATE OR REPLACE TABLE {namespace}.gold_product_sales
    USING DELTA AS
    SELECT
      stock_code,
      SUM(quantity) AS net_units,
      SUM(line_amount) AS net_signed_value_gbp
    FROM {namespace}.silver_order_items
    GROUP BY stock_code
    """)

    history = spark.table(
        table(config, "silver_customers_history")
    )

    if (
        history.filter("is_current")
        .groupBy("customer_id").count()
        .filter("count <> 1").limit(1).count()
    ):
        raise RuntimeError("SCD2 current-version check failed")

    silver_count = items.count()
    gold_count = (
        spark.table(table(config, "gold_daily_sales"))
        .agg(F.sum("line_count").alias("rows"))
        .first()["rows"] or 0
    )

    if silver_count != gold_count:
        raise RuntimeError("Silver-to-Gold count reconciliation failed")

    return {
        "silver_rows": silver_count,
        "gold_reconciled_rows": gold_count,
    }